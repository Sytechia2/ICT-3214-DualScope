#!/usr/bin/env python3
"""Build the reviewer pack for Task 9.5 from the selected sample.

For each sampled incident it writes one plain HTML sheet showing exactly the
evidence package the LLM sees (dualscope.investigation.package.build_package),
and three identical label workbooks (reviewer A, reviewer B, consensus).
The pack never contains LLM output, the Task 6.1 ATT&CK candidates or the
answer key; every package is checked with ``assert_no_answer_key`` and every
written file is scanned before the script finishes.

Outputs (in --output-dir, not committed):
  index.html                     instructions and the list of incidents
  incidents/<incident_id>.html   one evidence sheet per incident
  labels_reviewer_A.xlsx, labels_reviewer_B.xlsx, labels_consensus.xlsx

Example:
  python scripts/select_review_sample.py
  python scripts/build_review_pack.py
"""

from __future__ import annotations

import argparse
from html import escape
import json
from pathlib import Path
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

from dualscope.investigation.package import assert_no_answer_key, build_package  # noqa: E402

DEFAULT_HANDOFF = REPO_ROOT / "outputs" / "handoff" / "final_test_alerts_v1"
DEFAULT_SAMPLE = REPO_ROOT / "data" / "reference" / "evaluation" / "investigation_review_sample_v1.json"
DEFAULT_OUTPUT = REPO_ROOT / "outputs" / "evaluation" / "review_sample_v1"

HEADER = ["incident_id", "reviewer", "technique_id", "judgement", "evidence_refs", "notes"]
JUDGEMENTS = ["Supported", "Uncertain", "No supported mapping"]
WORKBOOKS = {"labels_reviewer_A.xlsx": "A", "labels_reviewer_B.xlsx": "B", "labels_consensus.xlsx": "consensus"}
FOOTER = ("Look up techniques at https://attack.mitre.org/. "
          "Do not open the dashboard's Investigation tab until your labels are submitted.")
FORBIDDEN_TEXT = re.compile(r"ground_truth|redteam|red-team|red_team", re.IGNORECASE)
TECHNIQUE_ID = re.compile(r"\bT1\d{3}")

INSTRUCTIONS = [
    "Label each incident independently, before you see any LLM output and before you open the dashboard's Investigation tab.",
    "Use only the incident sheet (incidents/<incident_id>.html) and https://attack.mitre.org/.",
    "One row per technique: technique_id (format T#### or T####.###), judgement, evidence_refs, notes. Add extra rows (copy incident_id) for more techniques.",
    "Supported = the cited events show the behaviour the technique describes (consistent with it, not proof of intent).",
    "Uncertain = plausible, but the logs cannot show it.",
    "If nothing fits, use one row with judgement 'No supported mapping' and leave technique_id empty.",
    "evidence_refs = semicolon-separated auth.txt references, e.g. auth.txt:123;auth.txt:456.",
    "Prefer the parent technique when the sub-technique cannot be told from logon records.",
    "The final labels go in labels_consensus.xlsx, the file the scores use.",
]

STYLE = """body{font-family:system-ui,-apple-system,"Segoe UI",sans-serif;font-size:14px;color:#222;margin:16px auto;max-width:1200px;padding:0 16px}
h1{font-size:18px}h2{font-size:15px;margin:18px 0 6px;border-bottom:1px solid #bbb}
table{border-collapse:collapse;font-size:13px;margin:4px 0}th,td{border:1px solid #bbb;padding:2px 6px;text-align:left;vertical-align:top}
th{background:#eee}.note{color:#666;font-size:12px}footer{margin-top:24px;padding-top:8px;border-top:1px solid #bbb;color:#555;font-size:13px}
pre{background:#f4f4f4;padding:6px;font-size:12px;overflow-x:auto}"""


def sheet_name(incident_id: str) -> str:
    """File name for an incident sheet (IDs can contain characters Windows rejects, e.g. '?')."""
    return re.sub(r"[^A-Za-z0-9._-]", "_", incident_id) + ".html"


def page(title: str, body: str) -> str:
    return (f"<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\"><title>{escape(title)}</title>"
            f"<style>{STYLE}</style></head><body>{body}</body></html>")


def table(headers: list[str], rows: list[list]) -> str:
    head = "".join(f"<th>{escape(str(h))}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in row) + "</tr>" for row in rows)
    return f"<table><tr>{head}</tr>{body}</table>"


def kv(mapping: dict) -> str:
    return table(["field", "value"], [[k, json.dumps(v) if isinstance(v, (dict, list)) else v]
                                       for k, v in mapping.items()])


def incident_page(package: dict) -> str:
    """The evidence package as a reviewer sheet; same content as the LLM prompt."""
    detector = package["detector"]
    counts = {k: v for k, v in package["counts"].items() if k != "note"}
    selection = package["selection"]
    parts = [f"<h1>{escape(package['incident_id'])}</h1>",
             "<p class=\"note\"><a href=\"../index.html\">Back to the list</a></p>",
             "<h2>Incident</h2>",
             kv({k: package[k] for k in ("incident_id", "user_id", "dataset_day", "start", "end", "duration_hours")}),
             "<h2>Detector</h2>", f"<p>Priority: <b>{escape(str(detector['priority']))}</b></p>"]
    hours = detector["alert_hours"]
    if hours:
        parts.append(table(list(hours[0]), [list(h.values()) for h in hours]))
    parts.append(f"<p class=\"note\">{escape(detector['notes'])}</p>")
    parts += ["<h2>Counts (all events of the incident)</h2>", kv(counts), f"<p class=\"note\">{escape(package['counts']['note'])}</p>"]
    graph = package["graph_context"]
    parts.append("<h2>Graph context (context only)</h2>")
    if graph:
        edges = graph["top_edges"]
        parts.append(kv({k: v for k, v in graph.items() if k not in ("top_edges", "notes")}))
        if edges:
            parts.append(table(list(edges[0]), [list(e.values()) for e in edges]))
        parts.append(f"<p class=\"note\">{escape(graph['notes'])}</p>")
    else:
        parts.append("<p>None for this incident.</p>")
    parts.append("<h2>Behaviours found by rules</h2>")
    if package["behaviours"]:
        parts.append(table(["behaviour", "description", "events", "example refs"],
                           [[b["behaviour"], b["description"], b["event_count"], "; ".join(b["example_refs"])]
                            for b in package["behaviours"]]))
    else:
        parts.append("<p>No rule matched any event.</p>")
    parts += [f"<h2>Events ({selection['shown_events']} of {selection['total_events']} shown)</h2>",
              f"<p class=\"note\">{escape(selection['rule'])}</p>"]
    parts.append(table(["ref", "time", "source user", "destination user", "source computer", "destination computer",
                        "auth type", "logon type", "orientation", "result", "flags"],
                       [[e["ref"], e["time"], e["source_user"], e["destination_user"], e["source_computer"],
                         e["destination_computer"], e["auth_type"], e["logon_type"], e["orientation"], e["result"],
                         "; ".join(e["flags"]) or "-"] for e in package["events"]]))
    parts.append(f"<footer>{escape(FOOTER)}</footer>")
    return page(package["incident_id"], "".join(parts))


def index_page(sample: dict, packages: dict[str, dict]) -> str:
    rows = [[f"<a href=\"incidents/{escape(sheet_name(i))}\">{escape(i)}</a>", packages[i]["detector"]["priority"],
             packages[i]["counts"]["events"]] for i in packages]
    body = "".join(f"<tr><td>{r[0]}</td><td>{escape(str(r[1]))}</td><td>{r[2]}</td></tr>" for r in rows)
    steps = "".join(f"<li>{escape(s)}</li>" for s in INSTRUCTIONS)
    return page("Review sample", f"<h1>Investigation review sample ({len(rows)} incidents)</h1>"
                f"<h2>Instructions</h2><ol>{steps}</ol><h2>Incidents</h2>"
                f"<table><tr><th>incident</th><th>priority</th><th>events</th></tr>{body}</table>"
                f"<footer>{escape(FOOTER)}</footer>")


def write_label_workbook(path: Path, reviewer: str, incident_ids: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "labels"
    sheet.append(HEADER)
    for incident_id in incident_ids:
        sheet.append([incident_id, reviewer, None, None, None, None])
    for column, width in zip("ABCDEF", (32, 12, 14, 22, 36, 50)):
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    validation = DataValidation(type="list", formula1='"' + ",".join(JUDGEMENTS) + '"', allow_blank=True)
    validation.add("D2:D500")
    sheet.add_data_validation(validation)
    guide = workbook.create_sheet("instructions")
    guide.column_dimensions["A"].width = 140
    for line in INSTRUCTIONS:
        guide.append([line])
        guide.cell(guide.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    listing = workbook.create_sheet("incidents")
    listing.append(["incident_id"])
    listing.column_dimensions["A"].width = 36
    for incident_id in incident_ids:
        listing.append([incident_id])
    workbook.save(path)


def load_events(handoff: Path, ids: list[str]) -> dict[str, pd.DataFrame]:
    alerts = pd.read_parquet(handoff / "alerts.parquet", columns=["alert_id", "incident_id"])
    events = pd.read_parquet(handoff / "events.parquet").merge(alerts, on="alert_id", validate="m:1")
    return dict(tuple(events[events["incident_id"].isin(ids)].groupby("incident_id", sort=True)))


def read_incidents(handoff: Path, ids: list[str]) -> dict[str, dict]:
    wanted = set(ids)
    found = {}
    for line in (handoff / "incidents.jsonl").read_text(encoding="utf-8").splitlines():
        if line.strip():
            record = json.loads(line)
            if record["incident_id"] in wanted:
                found[record["incident_id"]] = record
    return found


def check_files(output_dir: Path) -> None:
    """Fail if any generated file mentions the answer key or an ATT&CK ID on an incident sheet."""
    for path in sorted(output_dir.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix == ".xlsx":
            from openpyxl import load_workbook
            texts = [str(c.value) for ws in load_workbook(path) for row in ws.iter_rows() for c in row if c.value]
            text = "\n".join(texts)
        else:
            text = path.read_text(encoding="utf-8")
        if FORBIDDEN_TEXT.search(text):
            raise AssertionError(f"{path} mentions the answer key")
        if path.parent.name == "incidents" and TECHNIQUE_ID.search(text):
            raise AssertionError(f"{path} contains an ATT&CK technique ID")


def build_pack(incidents: dict[str, dict], events: dict[str, pd.DataFrame], sample: dict, output_dir: Path) -> list[str]:
    ids = sorted(entry["incident_id"] for entry in sample["incidents"])
    if len({sheet_name(i) for i in ids}) != len(ids):
        raise ValueError("incident IDs collide as file names")
    packages = {}
    for incident_id in ids:
        packages[incident_id] = build_package(incidents[incident_id], events[incident_id])
        assert_no_answer_key(packages[incident_id])
    (output_dir / "incidents").mkdir(parents=True, exist_ok=True)
    for incident_id, package in packages.items():
        (output_dir / "incidents" / sheet_name(incident_id)).write_text(incident_page(package), encoding="utf-8")
    (output_dir / "index.html").write_text(index_page(sample, packages), encoding="utf-8")
    for name, reviewer in WORKBOOKS.items():
        write_label_workbook(output_dir / name, reviewer, ids)
    check_files(output_dir)
    return ids


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--handoff", type=Path, default=DEFAULT_HANDOFF)
    parser.add_argument("--sample", type=Path, default=DEFAULT_SAMPLE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    sample = json.loads(args.sample.read_text(encoding="utf-8"))
    ids = [entry["incident_id"] for entry in sample["incidents"]]
    incidents = read_incidents(args.handoff, ids)
    events = load_events(args.handoff, ids)
    missing = sorted(set(ids) - set(incidents))
    if missing or set(ids) - set(events):
        raise SystemExit(f"sampled incidents not found in the handoff: {missing or sorted(set(ids) - set(events))}")
    built = build_pack(incidents, events, sample, args.output_dir)
    print(f"wrote {len(built)} incident sheets, index.html and {len(WORKBOOKS)} label workbooks to {args.output_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
