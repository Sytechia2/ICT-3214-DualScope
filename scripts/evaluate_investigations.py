#!/usr/bin/env python3
"""Score the three investigation modes against reviewer labels (Task 9.5).

Reviewers label the reference ATT&CK techniques for a sample of incidents
before they see any LLM output. This script compares what each mode mapped
(direct, rag, rag_verified) with those labels. Everything is computed only
over the sample incidents, and every rate is reported with numerator and
denominator.

Inputs:
  --run-dir          generated_direct.jsonl, generated_rag.jsonl, rag_verified.jsonl
  --sample           {"incidents": [{"incident_id", "strata"}]}
  --labels           consensus labels (.xlsx sheet "labels" or .csv) with columns
                     incident_id | reviewer | technique_id | judgement | evidence_refs | notes
                     judgement: Supported | Uncertain | No supported mapping
  --reviewer-labels  optional per-reviewer files (two) for agreement

Metric definitions (per mode; S = reference Supported IDs, U = reference
Uncertain IDs; "parent" = ID before the dot, so T1021.002 -> T1021):
  valid_id           mapped IDs found in the ATT&CK catalogue / all mapped
  precision_strict   mapped IDs in S / all mapped
  precision_parent   mapped IDs whose parent equals the parent of an S ID / all mapped
  uncertain_matches  mapped IDs in U (exact); neither correct nor unsupported
  unsupported        mapped IDs whose parent matches no S or U ID / all mapped
  recall_parent      S IDs whose parent was mapped / all S IDs
  empty_outputs      incidents with failed, invalid or missing replies / sample
  abstentions        ok replies without techniques; correct when the reference
                     has no Supported ID, missed when it has one
  false_mappings     no-mapping incidents where the mode mapped >= 1 technique
  verification       rag -> rag_verified: removed techniques (correct vs wrongly
                     removed, parent-level against S) and verifier status vs
                     reviewer judgement (exact ID match)

Outputs (in --output-dir): scores.json, per_incident.csv, report.md

Example:
  python scripts/evaluate_investigations.py --labels labels_consensus.xlsx
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys
from typing import Any, Iterable, Mapping

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

from dualscope.attack.catalog import TECHNIQUE_ID, TechniqueCatalog  # noqa: E402
from dualscope.attack.retrieval import DEFAULT_CATALOG  # noqa: E402

DEFAULT_RUN_DIR = REPO_ROOT / "outputs" / "investigations" / "gemini_v1"
DEFAULT_SAMPLE = REPO_ROOT / "data" / "reference" / "evaluation" / "investigation_review_sample_v1.json"
DEFAULT_LABELS = REPO_ROOT / "outputs" / "evaluation" / "review_sample_v1" / "labels_consensus.xlsx"
DEFAULT_OUTPUT = REPO_ROOT / "outputs" / "evaluation" / "investigation_scores_v1"

MODES = ("direct", "rag", "rag_verified")
COLUMNS = ("incident_id", "reviewer", "technique_id", "judgement", "evidence_refs", "notes")
SUPPORTED, UNCERTAIN, NO_MAPPING = "Supported", "Uncertain", "No supported mapping"
JUDGEMENTS = (SUPPORTED, UNCERTAIN, NO_MAPPING)


class LabelError(ValueError):
    """Raised when label files do not satisfy the labelling format."""


def parent(technique_id: str) -> str:
    return technique_id.split(".")[0]


def ratio(numerator: int, denominator: int) -> dict[str, Any]:
    return {"numerator": numerator, "denominator": denominator,
            "rate": numerator / denominator if denominator else None}


# ----------------------------------------------------------------------------
# Labels
# ----------------------------------------------------------------------------

def _cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def read_label_rows(path: Path) -> list[dict[str, str]]:
    """Rows of a label file (.xlsx sheet "labels" or .csv), values stripped."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            table = [list(row) for row in csv.reader(handle)]
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        if "labels" not in workbook.sheetnames:
            raise LabelError(f"{path}: no sheet named 'labels'")
        table = [list(row) for row in workbook["labels"].iter_rows(values_only=True)]
        workbook.close()
    if not table:
        raise LabelError(f"{path}: empty label file")
    header = [_cell(value) for value in table[0]]
    if tuple(header[:len(COLUMNS)]) != COLUMNS:
        raise LabelError(f"{path}: header must be {' | '.join(COLUMNS)}, got {' | '.join(header)}")
    rows = []
    for raw in table[1:]:
        values = [_cell(value) for value in raw] + [""] * len(COLUMNS)
        row = dict(zip(COLUMNS, values[:len(COLUMNS)]))
        if any(row.values()):
            rows.append(row)
    return rows


def build_reference(rows: Iterable[Mapping[str, str]], sample_ids: Iterable[str],
                    source: str = "labels") -> dict[str, dict[str, Any]]:
    """Validate label rows; return {incident_id: {"S": set, "U": set, "no_mapping": bool}}."""
    sample = list(sample_ids)
    known = set(sample)
    errors: list[str] = []
    reference: dict[str, dict[str, Any]] = {}
    for number, row in enumerate(rows, start=2):
        incident = row["incident_id"]
        judgement = row["judgement"]
        technique = row["technique_id"].upper()
        where = f"{source} row {number}"
        if incident not in known:
            errors.append(f"{where}: unknown incident {incident!r}")
            continue
        if judgement not in JUDGEMENTS:
            errors.append(f"{where}: bad judgement {judgement!r}")
            continue
        entry = reference.setdefault(incident, {"S": set(), "U": set(), "no_mapping": False})
        if judgement == NO_MAPPING:
            if technique:
                errors.append(f"{where}: 'No supported mapping' row must not have a technique ID")
            entry["no_mapping"] = True
            continue
        if not TECHNIQUE_ID.fullmatch(technique):
            errors.append(f"{where}: malformed technique ID {row['technique_id']!r}")
            continue
        entry["S" if judgement == SUPPORTED else "U"].add(technique)
    for incident, entry in reference.items():
        if entry["no_mapping"] and (entry["S"] or entry["U"]):
            errors.append(f"{source}: {incident} mixes 'No supported mapping' with technique rows")
        both = entry["S"] & entry["U"]
        if both:
            errors.append(f"{source}: {incident} labels {sorted(both)} both Supported and Uncertain")
    missing = [incident for incident in sample if incident not in reference]
    if missing:
        errors.append(f"{source}: {len(missing)} sample incident(s) have no label row: {', '.join(missing)}")
    if errors:
        raise LabelError("\n".join(errors))
    return reference


# ----------------------------------------------------------------------------
# Run outputs
# ----------------------------------------------------------------------------

def read_jsonl(path: Path) -> list[dict[str, Any]]:
    path = Path(path)
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _ids(items: Iterable[Mapping[str, Any]]) -> list[str]:
    seen: list[str] = []
    for item in items:
        technique_id = str(item.get("technique_id", "")).strip().upper()
        if technique_id and technique_id not in seen:
            seen.append(technique_id)
    return seen


def load_modes(run_dir: Path) -> dict[str, dict[str, dict[str, Any]]]:
    """{mode: {incident_id: {"ok", "mapped", ...}}} from a run folder."""
    run_dir = Path(run_dir)
    modes: dict[str, dict[str, dict[str, Any]]] = {mode: {} for mode in MODES}
    for mode in ("direct", "rag"):
        for record in read_jsonl(run_dir / f"generated_{mode}.jsonl"):
            ok = record.get("status") == "ok"
            techniques = ((record.get("response") or {}).get("techniques") or []) if ok else []
            modes[mode][record["incident_id"]] = {"ok": ok, "mapped": _ids(techniques)}
    for record in read_jsonl(run_dir / "rag_verified.jsonl"):
        ok = record.get("status") == "ok"
        verified = (record.get("verified") or {}) if ok else {}
        kept = verified.get("techniques") or []
        modes["rag_verified"][record["incident_id"]] = {
            "ok": ok,
            "mapped": _ids(kept),
            "kept": [{"technique_id": str(t.get("technique_id", "")).strip().upper(), "status": t.get("status")}
                     for t in kept],
            "removed": _ids(verified.get("removed_techniques") or []),
        }
    return modes


# ----------------------------------------------------------------------------
# Metrics
# ----------------------------------------------------------------------------

def reference_label(entry: Mapping[str, Any]) -> str:
    return "mapped" if (entry["S"] or entry["U"]) else "no_mapping"


def judge_id(technique_id: str, entry: Mapping[str, Any], catalog_ids: set[str]) -> str:
    """Verdict for one mapped ID against one incident's reference."""
    if technique_id not in catalog_ids:
        return "invalid_id"
    if technique_id in entry["S"]:
        return "supported"
    if technique_id in entry["U"]:
        return "uncertain"
    if parent(technique_id) in {parent(i) for i in entry["S"]}:
        return "supported_parent"
    if parent(technique_id) in {parent(i) for i in entry["U"]}:
        return "uncertain_parent"
    return "unsupported"


def score_mode(records: Mapping[str, Mapping[str, Any]], reference: Mapping[str, Mapping[str, Any]],
               catalog_ids: set[str]) -> dict[str, Any]:
    sample = list(reference)
    mapped_total = valid = strict = parent_ok = uncertain = unsupported = 0
    s_total = s_found = 0
    empty = abstain_correct = abstain_missed = false_maps = no_map_incidents = 0
    for incident in sample:
        entry = reference[incident]
        record = records.get(incident)
        no_mapping = reference_label(entry) == "no_mapping"
        no_map_incidents += no_mapping
        mapped = record["mapped"] if record and record["ok"] else []
        if not record or not record["ok"]:
            empty += 1
        elif not mapped:
            if entry["S"]:
                abstain_missed += 1
            else:
                abstain_correct += 1
        if no_mapping and mapped:
            false_maps += 1
        for technique in mapped:
            verdict = judge_id(technique, entry, catalog_ids)
            mapped_total += 1
            valid += technique in catalog_ids
            strict += verdict == "supported"
            parent_ok += verdict in ("supported", "supported_parent")
            uncertain += verdict == "uncertain"
            unsupported += verdict in ("unsupported", "invalid_id")
        mapped_parents = {parent(t) for t in mapped}
        s_total += len(entry["S"])
        s_found += sum(parent(i) in mapped_parents for i in entry["S"])
    abstained = abstain_correct + abstain_missed
    return {
        "valid_id_rate": ratio(valid, mapped_total),
        "precision_strict": ratio(strict, mapped_total),
        "precision_parent": ratio(parent_ok, mapped_total),
        "uncertain_matches": ratio(uncertain, mapped_total),
        "unsupported_rate": ratio(unsupported, mapped_total),
        "recall_parent": ratio(s_found, s_total),
        "empty_outputs": ratio(empty, len(sample)),
        "abstentions": {"total": abstained,
                        "correct": ratio(abstain_correct, abstained),
                        "missed": ratio(abstain_missed, abstained)},
        "false_mappings_on_no_mapping": ratio(false_maps, no_map_incidents),
        "mapped_techniques": mapped_total,
    }


def verification_impact(modes: Mapping[str, Mapping[str, Mapping[str, Any]]],
                        reference: Mapping[str, Mapping[str, Any]]) -> dict[str, Any]:
    """Effect of verification on the rag replies, judged against the reviewer labels."""
    correct = wrong = incidents = 0
    confusion = {verifier: {SUPPORTED: 0, UNCERTAIN: 0, "none": 0} for verifier in (SUPPORTED, UNCERTAIN)}
    for incident, entry in reference.items():
        record = modes["rag_verified"].get(incident)
        if not record or not record["ok"]:
            continue
        incidents += 1
        parents_s = {parent(i) for i in entry["S"]}
        for technique in record["removed"]:
            if parent(technique) in parents_s:
                wrong += 1
            else:
                correct += 1
        for kept in record["kept"]:
            verifier = kept["status"] if kept["status"] in confusion else UNCERTAIN
            if kept["technique_id"] in entry["S"]:
                reviewer = SUPPORTED
            elif kept["technique_id"] in entry["U"]:
                reviewer = UNCERTAIN
            else:
                reviewer = "none"
            confusion[verifier][reviewer] += 1
    removed = correct + wrong
    return {
        "incidents": incidents,
        "removed_total": removed,
        "correctly_removed": ratio(correct, removed),
        "supported_wrongly_removed": ratio(wrong, removed),
        "kept_confusion_verifier_vs_reviewer": confusion,
    }


def reviewer_agreement(reference_a: Mapping[str, Mapping[str, Any]],
                       reference_b: Mapping[str, Mapping[str, Any]]) -> dict[str, Any]:
    incidents = sorted(reference_a)
    jaccards, any_agree, disagreements = [], 0, 0
    for incident in incidents:
        a, b = reference_a[incident], reference_b[incident]
        pa = {parent(i) for i in a["S"]}
        pb = {parent(i) for i in b["S"]}
        jaccard = len(pa & pb) / len(pa | pb) if pa | pb else 1.0
        jaccards.append(jaccard)
        same_any = reference_label(a) == reference_label(b)
        any_agree += same_any
        disagreements += (jaccard < 1.0) or not same_any
    return {
        "incidents": len(incidents),
        "mean_jaccard_supported_parent": sum(jaccards) / len(jaccards) if jaccards else None,
        "any_mapping_agreement": ratio(any_agree, len(incidents)),
        "disagreeing_incidents": ratio(disagreements, len(incidents)),
    }


def per_incident_rows(modes, reference, catalog_ids) -> list[dict[str, str]]:
    rows = []
    for incident, entry in reference.items():
        row = {"incident_id": incident,
               "reference_supported": ";".join(sorted(entry["S"])),
               "reference_uncertain": ";".join(sorted(entry["U"])),
               "reference_kind": reference_label(entry)}
        for mode in MODES:
            record = modes[mode].get(incident)
            row[f"{mode}_status"] = "missing" if not record else ("ok" if record["ok"] else "failed")
            mapped = record["mapped"] if record and record["ok"] else []
            row[f"{mode}_mapped"] = ";".join(mapped)
            row[f"{mode}_verdicts"] = ";".join(f"{t}={judge_id(t, entry, catalog_ids)}" for t in mapped)
        rows.append(row)
    return rows


# ----------------------------------------------------------------------------
# Report
# ----------------------------------------------------------------------------

def fmt(value: Mapping[str, Any]) -> str:
    if value["rate"] is None:
        return f"n/a ({value['numerator']}/{value['denominator']})"
    return f"{value['rate']:.1%} ({value['numerator']}/{value['denominator']})"


def render_report(scores: Mapping[str, Any]) -> str:
    modes = scores["modes"]
    sample = scores["sample"]
    head = "| Metric | " + " | ".join(MODES) + " |"
    rule = "|---|" + "---|" * len(MODES)
    lines = [
        "# Investigation scores",
        "",
        f"Sample: {sample['incidents']} incidents ({sample['reference_with_mapping']} with a reference mapping, "
        f"{sample['reference_no_mapping']} with no supported mapping). Reference techniques were labelled by "
        "reviewers before they saw any LLM output. Rates are numerator/denominator.",
        "",
        "## Mapping quality",
        "",
        head,
        rule,
    ]
    for key, label in (("valid_id_rate", "Valid ATT&CK ID"), ("precision_strict", "Precision (strict)"),
                       ("precision_parent", "Precision (parent-level)"),
                       ("uncertain_matches", "Uncertain matches"), ("unsupported_rate", "Unsupported generation"),
                       ("recall_parent", "Recall (parent-level)")):
        lines.append(f"| {label} | " + " | ".join(fmt(modes[m][key]) for m in MODES) + " |")
    lines += ["", "## Empty outputs, abstentions and false mappings", "", head, rule]
    for key, label in (("empty_outputs", "Empty outputs (failed/invalid/missing)"),
                       ("false_mappings_on_no_mapping", "False mappings on no-mapping incidents")):
        lines.append(f"| {label} | " + " | ".join(fmt(modes[m][key]) for m in MODES) + " |")
    lines.append("| Correct abstentions | " + " | ".join(fmt(modes[m]["abstentions"]["correct"]) for m in MODES) + " |")
    lines.append("| Missed abstentions (reference has Supported) | "
                 + " | ".join(fmt(modes[m]["abstentions"]["missed"]) for m in MODES) + " |")
    v = scores["verification_impact"]
    lines += ["", "## Verification impact (rag -> rag_verified)", "",
              "| Metric | Value |", "|---|---|",
              f"| Techniques removed | {v['removed_total']} |",
              f"| Correctly removed (parent not Supported) | {fmt(v['correctly_removed'])} |",
              f"| Supported mappings wrongly removed | {fmt(v['supported_wrongly_removed'])} |",
              "", "Kept techniques: verifier status vs reviewer judgement (exact ID).", "",
              "| Verifier / Reviewer | Supported | Uncertain | No label |", "|---|---|---|---|"]
    for verifier, counts in v["kept_confusion_verifier_vs_reviewer"].items():
        lines.append(f"| {verifier} | {counts[SUPPORTED]} | {counts[UNCERTAIN]} | {counts['none']} |")
    agreement = scores.get("reviewer_agreement")
    if agreement:
        lines += ["", "## Reviewer agreement", "", "| Metric | Value |", "|---|---|",
                  f"| Mean Jaccard of Supported sets (parent-level) | {agreement['mean_jaccard_supported_parent']:.3f} |",
                  f"| Agreement on any mapping vs none | {fmt(agreement['any_mapping_agreement'])} |",
                  f"| Incidents with any disagreement | {fmt(agreement['disagreeing_incidents'])} |"]
    return "\n".join(lines) + "\n"


# ----------------------------------------------------------------------------
# Driver
# ----------------------------------------------------------------------------

def sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_sample_ids(path: Path) -> list[str]:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    ids = [item["incident_id"] for item in data["incidents"]]
    if len(ids) != len(set(ids)):
        raise LabelError(f"{path}: duplicate incident IDs in sample")
    return ids


def evaluate(run_dir: Path, sample_path: Path, labels_path: Path,
             reviewer_paths: Iterable[Path] = (), catalog_path: Path = DEFAULT_CATALOG) -> dict[str, Any]:
    reviewer_paths = list(reviewer_paths)
    sample_ids = load_sample_ids(sample_path)
    reference = build_reference(read_label_rows(labels_path), sample_ids, source=Path(labels_path).name)
    reference = {incident: reference[incident] for incident in sample_ids}
    catalog_ids = {t.technique_id for t in TechniqueCatalog.load(catalog_path).techniques}
    modes = load_modes(run_dir)
    scores: dict[str, Any] = {
        "created_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "sample": {"incidents": len(sample_ids),
                   "reference_with_mapping": sum(reference_label(e) == "mapped" for e in reference.values()),
                   "reference_no_mapping": sum(reference_label(e) == "no_mapping" for e in reference.values()),
                   "reference_supported_ids": sum(len(e["S"]) for e in reference.values()),
                   "reference_uncertain_ids": sum(len(e["U"]) for e in reference.values())},
        "modes": {mode: score_mode(modes[mode], reference, catalog_ids) for mode in MODES},
        "verification_impact": verification_impact(modes, reference),
    }
    inputs = {"sample": sample_path, "labels": labels_path, "catalog": catalog_path}
    for mode_file in ("generated_direct.jsonl", "generated_rag.jsonl", "rag_verified.jsonl"):
        if (Path(run_dir) / mode_file).exists():
            inputs[mode_file] = Path(run_dir) / mode_file
    for index, path in enumerate(reviewer_paths):
        inputs[f"reviewer_labels_{index + 1}"] = path
    scores["inputs"] = {name: {"path": str(path), "sha256": sha256(path)} for name, path in inputs.items()}
    if len(reviewer_paths) >= 2:
        refs = [build_reference(read_label_rows(p), sample_ids, source=Path(p).name) for p in reviewer_paths[:2]]
        scores["reviewer_agreement"] = reviewer_agreement(*refs)
    scores["_per_incident"] = per_incident_rows(modes, reference, catalog_ids)
    return scores


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--run-dir", type=Path, default=DEFAULT_RUN_DIR)
    parser.add_argument("--sample", type=Path, default=DEFAULT_SAMPLE)
    parser.add_argument("--labels", type=Path, default=DEFAULT_LABELS)
    parser.add_argument("--reviewer-labels", type=Path, nargs="+", default=[])
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    if len(args.reviewer_labels) == 1:
        parser.error("--reviewer-labels needs two files")
    try:
        scores = evaluate(args.run_dir, args.sample, args.labels, args.reviewer_labels, args.catalog)
    except LabelError as error:
        print(f"Label validation failed:\n{error}", file=sys.stderr)
        return 1
    rows = scores.pop("_per_incident")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "scores.json").write_text(json.dumps(scores, indent=2) + "\n", encoding="utf-8")
    with (args.output_dir / "per_incident.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    report = render_report(scores)
    (args.output_dir / "report.md").write_text(report, encoding="utf-8")
    print(report)
    print(f"Wrote {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
