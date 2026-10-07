import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from scripts.build_review_pack import HEADER, JUDGEMENTS, TECHNIQUE_ID, build_pack, check_files, sheet_name
from scripts.select_review_sample import incident_strata, select

DAY = 86_400


def _event(line, ts, **overrides):
    event = {
        "source_reference": f"auth.txt:{line}", "source_line": line, "timestamp": ts,
        "acting_user": "U1@DOM1", "source_user": "U1@DOM1", "destination_user": "U1@DOM1",
        "source_computer": "C1", "destination_computer": "C2", "authentication_type": "Kerberos",
        "logon_type": "Network", "authentication_orientation": "LogOn", "authentication_result": "Success",
        "is_new_user_source": False, "is_new_host_connection": False, "is_new_user_destination": False,
        "prior_auth_count_1h": 0, "prior_failure_count_1h": 0, "prior_unique_destinations_24h": 1,
    }
    event.update(overrides)
    return event


def _world():
    """12 incidents over 6 days; some with failures, new destinations, graph agreement, one red-team."""
    incidents, events = {}, {}
    line = 0
    for n in range(12):
        day = 17 + n % 6
        iid = f"INC-TEST-D{day}-U{n}_DOM1-001"
        start = (day - 1) * DAY + 3_600
        incidents[iid] = {
            "incident_id": iid, "user_id": f"U{n}@DOM1", "dataset_day": day, "start_time": start,
            "end_time": start + 3_600, "duration_hours": 1, "priority": "MEDIUM" if n % 3 == 0 else "HIGH",
            "ground_truth_redteam": n == 5, "ground_truth_redteam_hours": int(n == 5),
            "alert_hours": [{"alert_id": f"A{n}", "window_start": start, "window_end": start + 3_600,
                             "fusion_score": 0.9, "rank_in_day": 1, "tied_at_cutoff": False,
                             "gru_percentile_in_day": 0.99, "ground_truth_redteam": n == 5}],
            "detector_scores": {"graph": {"max_score": 0.5, "any_alert": n % 4 == 0}},
            "graph_context": None,
        }
        rows = []
        for k in range(3 + n):
            line += 1
            extra = {}
            if n % 2 == 0 and k == 0:
                extra = {"authentication_result": "Fail"}
            if n % 3 == 1 and k == 0:
                extra = {"is_new_user_destination": True}
            rows.append(_event(line, start + k, **extra))
        events[iid] = pd.DataFrame(rows).assign(alert_id=f"A{n}", incident_id=iid)
    return incidents, events


TARGETS = {"priority_medium": 2, "graph_agrees": 2, "failed_logons": 2, "network_logon_to_new_destination": 2,
           "no_rule_match": 1}


@pytest.fixture()
def world():
    incidents, events = _world()
    strata = {i: incident_strata(incidents[i], events[i]) for i in incidents}
    forced = [i for i in incidents if incidents[i]["ground_truth_redteam"]]
    return incidents, events, strata, forced


def test_selection_is_deterministic_and_meets_targets(world):
    _, _, strata, forced = world
    first = select(strata, forced, 8, TARGETS, 2, seed=0)
    assert first == select(strata, forced, 8, TARGETS, 2, seed=0)
    assert first == sorted(first) and len(first) == 8
    assert set(forced) <= set(first)
    from collections import Counter
    assert max(Counter(strata[i]["dataset_day"] for i in first).values()) <= 2


def test_unmeetable_targets_raise(world):
    _, _, strata, forced = world
    with pytest.raises(ValueError):
        select(strata, forced, 8, {"priority_medium": 50}, 2)


def test_strata_do_not_reveal_the_forced_incident(world):
    _, _, strata, _ = world
    assert "ground_truth" not in json.dumps(strata) and "redteam" not in json.dumps(strata)


def test_pack_has_no_answer_key_or_attack_ids_and_valid_workbooks(world, tmp_path):
    incidents, events, strata, forced = world
    ids = select(strata, forced, 8, TARGETS, 2)
    sample = {"incidents": [{"incident_id": i, "strata": strata[i]} for i in ids]}
    built = build_pack(incidents, events, sample, tmp_path)
    assert built == ids
    for incident_id in ids:
        text = (tmp_path / "incidents" / sheet_name(incident_id)).read_text(encoding="utf-8")
        assert not TECHNIQUE_ID.search(text)
        assert "ground_truth" not in text and "redteam" not in text.lower()
        assert "auth.txt:" in text and "attack.mitre.org" in text
    for name, reviewer in (("labels_reviewer_A.xlsx", "A"), ("labels_reviewer_B.xlsx", "B"),
                           ("labels_consensus.xlsx", "consensus")):
        workbook = load_workbook(tmp_path / name)
        assert workbook.sheetnames == ["labels", "instructions", "incidents"]
        sheet = workbook["labels"]
        assert [c.value for c in sheet[1]] == HEADER
        assert [r[0].value for r in sheet.iter_rows(min_row=2)] == ids
        assert {r[1].value for r in sheet.iter_rows(min_row=2)} == {reviewer}
        assert sheet.freeze_panes == "A2"
        validations = sheet.data_validations.dataValidation
        assert len(validations) == 1 and str(validations[0].sqref).startswith("D2")
        assert validations[0].formula1 == '"' + ",".join(JUDGEMENTS) + '"'
    check_files(tmp_path)


def test_check_files_catches_leaks(tmp_path):
    (tmp_path / "incidents").mkdir()
    (tmp_path / "incidents" / "x.html").write_text("candidate T1021", encoding="utf-8")
    with pytest.raises(AssertionError):
        check_files(tmp_path)
